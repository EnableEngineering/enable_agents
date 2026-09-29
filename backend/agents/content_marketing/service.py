"""
Content Marketing Agent — service layer.

Uses SQLAlchemy models for all DB operations.
All SQLite code has been migrated to PostgreSQL.
"""
from datetime import datetime
from uuid import uuid4
import json
import os
from typing import Any, Dict, List, Optional

from flask import g, jsonify, request
from werkzeug.utils import secure_filename

from core.auth import user_can_access_project
from core.database import db
from .models import CMProject, CMDocument, CMKnowledgeGraph, CMGeneratedContent, CMConversation

CONTENT_MARKETING_UPLOAD_FOLDER = os.environ.get(
    "CONTENT_MARKETING_UPLOAD_FOLDER",
    os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
        "data",
        "content_marketing_uploads",
    ),
)
CONTENT_MARKETING_ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "xlsx", "html", "md"}
os.makedirs(CONTENT_MARKETING_UPLOAD_FOLDER, exist_ok=True)


def _extract_content(file_path: str, file_type: str) -> str:
    if file_type == "pdf":
        import fitz

        with fitz.open(file_path) as document:
            return "\n".join(page.get_text() for page in document)
    if file_type == "docx":
        from docx import Document

        document = Document(file_path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if file_type in {"txt", "md"}:
        with open(file_path, "r", encoding="utf-8") as source_file:
            return source_file.read()
    if file_type == "html":
        from bs4 import BeautifulSoup

        with open(file_path, "r", encoding="utf-8") as source_file:
            return BeautifulSoup(source_file.read(), "html.parser").get_text(" ")
    if file_type == "xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            return "\n".join(
                "\t".join(str(value) for value in row if value is not None)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
                if any(value is not None for value in row)
            )
        finally:
            workbook.close()
    return ""


def _analyze_domain(
    documents: List[str],
    user_id: str,
    platform_project_id: Optional[str],
    fallback_industry: Optional[str] = None,
) -> Dict[str, Any]:
    fallback = {
        "industry": fallback_industry or "General",
        "sector": "Unknown",
        "function": "Marketing",
        "role": "Marketing Manager",
        "target_audience": "Business Professionals",
        "value_proposition": "",
        "tone": "professional",
        "key_themes": [],
    }
    try:
        from core.ai_client import get_langchain_llm, log_langchain_usage

        llm, key_source, model = get_langchain_llm(
            user_id, platform_project_id, model="gpt-4", temperature=0
        )
        prompt = f"""Analyze these business documents and return one JSON object with
industry, sector, function, role, target_audience, value_proposition, tone,
and key_themes (an array of strings). Use concise values.

Documents:
{' '.join(documents[:3])[:2000]}"""
        response = llm.invoke(prompt)
        log_langchain_usage(
            response, user_id, platform_project_id,
            "content_marketing.analyze_documents", model, key_source,
        )
        response_text = str(response.content).strip()
        start, end = response_text.find("{"), response_text.rfind("}")
        if start >= 0 and end > start:
            result = json.loads(response_text[start:end + 1])
            return {**fallback, **result}
    except Exception:
        pass
    return fallback


def _owned_cm_project_or_none(project_id: str):
    """A CMProject only counts as accessible if it belongs to the caller -
    otherwise any authenticated user could pass another user's project_id
    and read/generate/chat against their documents."""
    if not project_id:
        return None
    project = CMProject.query.filter_by(project_id=project_id).first()
    if not project or project.user_id != g.user_id:
        return None
    return project


# =============================================================================
# Project Operations
# =============================================================================

def create_project():
    """Create a new content marketing project.

    If a `platform_project_id` is supplied (the platform-wide Project the
    user picked via the header ProjectSelector), this is idempotent: the
    content-marketing project is derived deterministically from it, so
    repeated calls (e.g. on every page load) reuse the same CMProject
    instead of piling up duplicates.
    """
    data = request.get_json(silent=True) or {}

    user_id = g.user_id
    project_name = data.get('project_name', 'Untitled Project')
    platform_project_id = data.get('platform_project_id')

    if platform_project_id and not user_can_access_project(user_id, platform_project_id):
        return jsonify({"success": False, "error": "Project not found"}), 404

    if platform_project_id:
        derived_id = f"cmp_{platform_project_id.replace('-', '')[:28]}"
        existing = CMProject.query.filter_by(project_id=derived_id).first()
        if existing:
            if existing.user_id != user_id:
                return jsonify({"success": False, "error": "Project not found"}), 404
            return jsonify({
                "success": True,
                "project_id": existing.project_id,
                "message": "Existing project reused",
            }), 200
        new_project_id = derived_id
    else:
        new_project_id = f"project_{uuid4().hex[:12]}"

    project = CMProject(
        project_id=new_project_id,
        user_id=user_id,
        platform_project_id=platform_project_id,
        project_name=project_name,
        description=data.get("description"),
        industry=data.get("industry"),
        sector=data.get("sector"),
        function=data.get("function"),
        role=data.get("role"),
    )
    db.session.add(project)
    try:
        db.session.commit()
    except Exception:
        # Two near-simultaneous requests (e.g. React effects firing twice)
        # can both pass the "existing" check above before either commits -
        # fall back to the row the other request just created instead of 500ing.
        db.session.rollback()
        existing = CMProject.query.filter_by(project_id=new_project_id).first()
        if existing:
            return jsonify({
                "success": True,
                "project_id": existing.project_id,
                "message": "Existing project reused",
            }), 200
        raise

    # Store in context lake
    try:
        from core.context import ContextStore
        ContextStore().set(
            user_id,
            "content_marketing",
            f"project:{project.project_id}",
            {
                "project_id": project.project_id,
                "project_name": project.project_name,
                "industry": project.industry,
                "updated_at": datetime.utcnow().isoformat(),
            },
        )
    except Exception:
        pass

    return jsonify({
        "success": True,
        "project_id": project.project_id,
        "message": f'Project "{project_name}" created successfully'
    }), 201


def get_project(project_id: str):
    """Get project details with statistics."""
    project = _owned_cm_project_or_none(project_id)
    if not project:
        return jsonify({"success": False, "error": "Project not found"}), 404

    doc_count = CMDocument.query.filter_by(project_id=project_id).count()
    kg = CMKnowledgeGraph.query.filter_by(project_id=project_id).first()

    return jsonify({
        "success": True,
        "project": project.to_dict(),
        "statistics": {
            "documents": doc_count,
            "has_knowledge_graph": kg is not None
        }
    })


# =============================================================================
# Document Operations
# =============================================================================

def upload_documents():
    project_id = request.form.get("project_id")
    if not project_id:
        return jsonify({"success": False, "error": "project_id required"}), 400

    uploaded_files = request.files.getlist("files")
    if not uploaded_files:
        return jsonify({"success": False, "error": "No files provided"}), 400

    project = _owned_cm_project_or_none(project_id)
    if not project:
        return jsonify({"success": False, "error": "Project not found"}), 404

    doc_ids = []
    document_texts = []
    try:
        for uploaded_file in uploaded_files:
            if not uploaded_file.filename:
                continue
            filename = secure_filename(uploaded_file.filename)
            if not filename or "." not in filename:
                continue
            file_type = filename.rsplit(".", 1)[-1].lower()
            if file_type not in CONTENT_MARKETING_ALLOWED_EXTENSIONS:
                continue

            safe_project_id = secure_filename(project_id)
            file_path = os.path.join(
                CONTENT_MARKETING_UPLOAD_FOLDER, safe_project_id, filename
            )
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            uploaded_file.save(file_path)
            extracted_content = _extract_content(file_path, file_type)
            doc_id = f"doc_{uuid4().hex[:12]}"
            db.session.add(
                CMDocument(
                    doc_id=doc_id,
                    project_id=project_id,
                    file_name=filename,
                    file_type=file_type,
                    file_path=file_path,
                    file_size=os.path.getsize(file_path),
                    document_type=file_type,
                    extracted_content=extracted_content,
                )
            )
            doc_ids.append(doc_id)
            document_texts.append(extracted_content)

        domain_context = _analyze_domain(
            document_texts, g.user_id, project.platform_project_id, project.industry
        )
        knowledge_graph_id = f"kg_{uuid4().hex[:12]}"
        entities = domain_context.get("key_themes") or []
        graph_data = {
            "entities": entities,
            "relationships": [],
            "domain_context": domain_context,
            "documents_count": len(doc_ids),
        }
        graph = CMKnowledgeGraph(kg_id=knowledge_graph_id, project_id=project_id)
        graph.kg_data = graph_data
        graph.entities = len(entities)
        graph.relationships = 0
        db.session.add(graph)
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        return jsonify({"success": False, "error": str(exc)}), 500

    return jsonify({
        "success": True,
        "uploaded_files": len(doc_ids),
        "document_ids": doc_ids,
        "knowledge_graph_id": knowledge_graph_id,
        "domain_specialization": domain_context,
    }), 201


def list_documents(project_id: str):
    """List all documents in a project."""
    if not _owned_cm_project_or_none(project_id):
        return jsonify({"success": False, "error": "Project not found"}), 404
    docs = CMDocument.query.filter_by(project_id=project_id).all()
    return jsonify({
        "success": True,
        "documents": [d.to_dict() for d in docs]
    })


def delete_document(doc_id: str):
    document = CMDocument.query.filter_by(doc_id=doc_id).first()
    if not document or not _owned_cm_project_or_none(document.project_id):
        return jsonify({"success": False, "error": "Document not found"}), 404
    try:
        if document.file_path and os.path.exists(document.file_path):
            os.remove(document.file_path)
        db.session.delete(document)
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        return jsonify({"success": False, "error": str(exc)}), 500
    return jsonify({"success": True}), 200


# =============================================================================
# Knowledge Graph
# =============================================================================

def get_knowledge_graph(project_id: str):
    """Retrieve knowledge graph for visualization."""
    if not _owned_cm_project_or_none(project_id):
        return jsonify({"success": False, "error": "Knowledge graph not found"}), 404
    kg = CMKnowledgeGraph.query.filter_by(project_id=project_id).first()
    if not kg:
        return jsonify({"success": False, "error": "Knowledge graph not found"}), 404

    return jsonify({
        "success": True,
        "kg_id": kg.kg_id,
        "kg_data": kg.kg_data,
        "entities": kg.entities,
        "relationships": kg.relationships,
        "created_at": kg.created_at.isoformat() if kg.created_at else None
    })


def generate_content():
    data = request.get_json(silent=True) or {}
    project_id = data.get("project_id")
    if not project_id:
        return jsonify({"success": False, "error": "project_id required"}), 400

    project = _owned_cm_project_or_none(project_id)
    if not project:
        return jsonify({"success": False, "error": "Project not found"}), 404

    documents = CMDocument.query.filter_by(project_id=project_id).all()
    document_texts = [doc.extracted_content for doc in documents if doc.extracted_content]
    if not document_texts:
        return jsonify({"success": False, "error": "No documents found in project"}), 400

    knowledge_graph = (
        CMKnowledgeGraph.query.filter_by(project_id=project_id)
        .order_by(CMKnowledgeGraph.created_at.desc())
        .first()
    )
    try:
        from core.settings import get_response_language_instruction
        from .rag_content_generator import RAGContentGenerator

        result = RAGContentGenerator(g.user_id, project.platform_project_id).generate(
            documents=document_texts,
            knowledge_graph=knowledge_graph.kg_data if knowledge_graph else None,
            channel=data.get("channel", "linkedin"),
            content_type=data.get("content_type", "post"),
            domain_context=project.industry,
            user_context=data.get("context", ""),
            language_instruction=get_response_language_instruction(g.user_id),
        )
        content_id = f"content_{uuid4().hex[:12]}"
        content = CMGeneratedContent(
            content_id=content_id,
            project_id=project_id,
            channel=result["metadata"]["channel"],
            content_type=result["metadata"]["content_type"],
            content=result["content"],
        )
        content.source_docs = [doc.doc_id for doc in documents]
        content.domain_context = {
            "industry": project.industry,
            "prompt": data.get("context", ""),
            "sources_used": result["metadata"]["sources_used"],
        }
        db.session.add(content)
        db.session.commit()
        return jsonify({"success": True, "content_id": content_id, **result}), 201
    except Exception as exc:
        db.session.rollback()
        return jsonify({"success": False, "error": str(exc)}), 500


def chat():
    data = request.get_json(silent=True) or {}
    project_id = data.get("project_id")
    message = data.get("message")
    if not project_id or not message:
        return jsonify({"success": False, "error": "project_id and message required"}), 400

    project = _owned_cm_project_or_none(project_id)
    if not project:
        return jsonify({"success": False, "error": "Project not found"}), 404

    documents = CMDocument.query.filter_by(project_id=project_id).limit(10).all()
    knowledge_graph = (
        CMKnowledgeGraph.query.filter_by(project_id=project_id)
        .order_by(CMKnowledgeGraph.created_at.desc())
        .first()
    )
    try:
        from core.settings import get_response_language_instruction
        from .rag_content_generator import RAGContentGenerator

        response = RAGContentGenerator(g.user_id, project.platform_project_id).chat_response(
            user_message=message,
            documents=[doc.extracted_content for doc in documents if doc.extracted_content],
            knowledge_graph=knowledge_graph.kg_data if knowledge_graph else None,
            language_instruction=get_response_language_instruction(g.user_id),
        )
        message_id = f"msg_{uuid4().hex[:12]}"
        conversation = CMConversation(
            msg_id=message_id,
            project_id=project_id,
            user_message=message,
            agent_response=response,
        )
        conversation.context = {"project_name": project.project_name, "doc_count": len(documents)}
        db.session.add(conversation)
        db.session.commit()
        return jsonify({"success": True, "response": response, "message_id": message_id}), 200
    except Exception as exc:
        db.session.rollback()
        return jsonify({"success": False, "error": str(exc)}), 500


# =============================================================================
# Plain-argument core (LangGraph orchestration)
# =============================================================================

_CHANNEL_CONFIG = {
    'linkedin': {'tone': 'professional', 'max_length': 3000},
    'email': {'tone': 'persuasive', 'max_length': 500},
    'social': {'tone': 'casual', 'max_length': 280},
    'google_ads': {'tone': 'direct', 'max_length': 150},
}


def generate_content_core(channel, content_type, user_context, user_id,
                           industry="General", doc_texts=None, cm_project_id=None,
                           source_doc_ids=None):
    """Plain-argument core of app.py's generate_content_marketing - callable
    from a LangGraph node (or anywhere else outside a Flask request) with no
    request/g dependency.

    Unlike the interactive route (which hard-requires an existing CMProject
    with uploaded CMDocuments), `doc_texts` is optional here - a workflow
    node has no human-uploaded documents to draw on, so this degrades to
    generating from `user_context`/`industry` alone. `cm_project_id` is also
    optional: CMGeneratedContent.project_id is a NOT NULL foreign key to
    cm_projects, so a generated-content row is only persisted when a real
    CMProject id is given; otherwise the content is returned but not saved
    anywhere, the same way graph.py's document_analysis_node returns its RAG
    answer without creating a permanent record.

    Returns (result_dict_or_None, error_message_or_None).
    """
    doc_texts = doc_texts or []
    config = _CHANNEL_CONFIG.get(channel, _CHANNEL_CONFIG['linkedin'])

    from core.settings import get_response_language_instruction
    prompt = f"""Generate marketing content for {channel} channel.
Industry: {industry or 'General'}
Tone: {config['tone']}
Max Length: {config['max_length']} characters
Content Type: {content_type}
User Context: {user_context}
Documents Summary: {' '.join([doc[:200] for doc in doc_texts[:3]])}

Language level: {get_response_language_instruction(user_id)}

Generate compelling marketing {content_type} content."""

    try:
        from core.ai_client import get_langchain_llm, log_langchain_usage

        ai_project_id = None
        if cm_project_id:
            project = CMProject.query.filter_by(project_id=cm_project_id).first()
            ai_project_id = project.platform_project_id if project else None

        llm, key_source, resolved_model = get_langchain_llm(user_id, ai_project_id, model="gpt-4", temperature=0.7)
        result = llm.invoke(prompt)
        log_langchain_usage(result, user_id, ai_project_id, "content_marketing.generate_content", resolved_model, key_source)
        response = result.content

        content_id = f"content_{uuid4().hex[:12]}"
        if cm_project_id and CMProject.query.filter_by(project_id=cm_project_id).first():
            content = CMGeneratedContent(
                content_id=content_id,
                project_id=cm_project_id,
                channel=channel,
                content_type=content_type,
                content=response,
            )
            content.source_docs = source_doc_ids or []
            content.domain_context = {"industry": industry, "prompt": user_context}
            db.session.add(content)
            db.session.commit()

        return {
            "content_id": content_id,
            "channel": channel,
            "content_type": content_type,
            "content": response,
            "variations": [response],
            "metadata": config,
        }, None
    except Exception as e:
        db.session.rollback()
        return None, str(e)
